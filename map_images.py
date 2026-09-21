# -*- coding: utf-8 -*-
import docx

doc = docx.Document('CARTILHA_Proposta 2.docx')
for i, p in enumerate(doc.paragraphs):
    if 'drawing' in p._element.xml:
        print(f'Paragraph {i} has image! Nearby text: "{p.text[:100]}"')
        for j in range(max(0, i-3), min(len(doc.paragraphs), i+4)):
            if doc.paragraphs[j].text.strip():
                print(f'   [{j}]: "{doc.paragraphs[j].text.strip()[:90]}"')
        print('-'*50)
